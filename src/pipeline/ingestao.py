import csv
import unicodedata
from pathlib import Path

import pandas as pd

from src.pipeline.config import INVENTARIO_ARQUIVOS, PROJECT_ROOT, RELATORIO_INGESTAO


COLUNAS_REFERENCIA = {
    "ano",
    "no_regiao",
    "sg_uf",
    "co_municipio",
    "no_municipio",
    "co_entidade",
    "no_entidade",
    "tipoloca",
    "dependad",
}

PREFIXOS_TAXAS = ("tap_", "tre_", "tab_")


def normalizar_texto(valor: object) -> str:
    if pd.isna(valor):
        return ""

    texto = str(valor).strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(caractere for caractere in texto if not unicodedata.combining(caractere))
    return texto


def caminho_absoluto(caminho_relativo: str) -> Path:
    return PROJECT_ROOT / caminho_relativo


def carregar_inventario(caminho_inventario: Path = INVENTARIO_ARQUIVOS) -> list[dict[str, str]]:
    if not caminho_inventario.exists():
        return []

    with caminho_inventario.open("r", newline="", encoding="utf-8") as arquivo_csv:
        return list(csv.DictReader(arquivo_csv))


def escolher_aba(abas: list[str]) -> tuple[str, str]:
    for aba in abas:
        if normalizar_texto(aba) == "escolas":
            return aba, "Aba ESCOLAS encontrada."

    for aba in abas:
        if "escola" in normalizar_texto(aba):
            return aba, "Aba ESCOLAS não encontrada; usada aba com termo escola."

    return abas[0], "Aba ESCOLAS não encontrada; usada a primeira aba disponível."


def pontuar_linha_cabecalho(valores: list[object]) -> int:
    textos = [normalizar_texto(valor) for valor in valores]
    textos = [texto for texto in textos if texto]

    pontuacao = 0
    for texto in textos:
        if texto in COLUNAS_REFERENCIA:
            pontuacao += 3
        if texto.startswith(PREFIXOS_TAXAS):
            pontuacao += 2
        if "entidade" in texto or "municipio" in texto:
            pontuacao += 1

    return pontuacao


def detectar_linha_cabecalho(caminho: Path, aba: str) -> tuple[int, str]:
    if caminho.suffix.lower() == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(caminho, on_demand=True)
        try:
            planilha = workbook.sheet_by_name(aba)
            melhor_linha = 0
            melhor_pontuacao = -1
            limite_linhas = min(30, planilha.nrows)

            for indice in range(limite_linhas):
                pontuacao = pontuar_linha_cabecalho(planilha.row_values(indice))
                if pontuacao > melhor_pontuacao:
                    melhor_linha = indice
                    melhor_pontuacao = pontuacao
        finally:
            workbook.release_resources()

        if melhor_pontuacao <= 0:
            return 0, "Cabeçalho não identificado automaticamente; usada a primeira linha."

        return melhor_linha, "Cabeçalho identificado automaticamente."

    previa = pd.read_excel(caminho, sheet_name=aba, header=None, nrows=30, dtype=str)

    melhor_linha = 0
    melhor_pontuacao = -1
    for indice, linha in previa.iterrows():
        pontuacao = pontuar_linha_cabecalho(linha.tolist())
        if pontuacao > melhor_pontuacao:
            melhor_linha = int(indice)
            melhor_pontuacao = pontuacao

    if melhor_pontuacao <= 0:
        return 0, "Cabeçalho não identificado automaticamente; usada a primeira linha."

    return melhor_linha, "Cabeçalho identificado automaticamente."


def detectar_dimensoes_planilha(caminho: Path, aba: str) -> tuple[int, int]:
    if caminho.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(caminho, read_only=True, data_only=True)
        try:
            planilha = workbook[aba]
            return planilha.max_row or 0, planilha.max_column or 0
        finally:
            workbook.close()

    if caminho.suffix.lower() == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(caminho, on_demand=True)
        try:
            planilha = workbook.sheet_by_name(aba)
            return planilha.nrows, planilha.ncols
        finally:
            workbook.release_resources()

    dados = pd.read_excel(caminho, sheet_name=aba, header=None, dtype=str)
    return len(dados), len(dados.columns)


def carregar_metadados_planilha(caminho: Path) -> dict[str, str]:
    arquivo_excel = pd.ExcelFile(caminho)
    aba, observacao_aba = escolher_aba(arquivo_excel.sheet_names)
    linha_cabecalho, observacao_cabecalho = detectar_linha_cabecalho(caminho, aba)
    total_linhas, total_colunas = detectar_dimensoes_planilha(caminho, aba)
    linhas_dados = max(total_linhas - linha_cabecalho - 1, 0)

    return {
        "aba_utilizada": aba,
        "linha_cabecalho": str(linha_cabecalho + 1),
        "quantidade_linhas": str(linhas_dados),
        "quantidade_colunas": str(total_colunas),
        "observacao": f"{observacao_aba} {observacao_cabecalho}",
    }


def carregar_planilha(caminho: Path) -> tuple[pd.DataFrame, dict[str, str]]:
    arquivo_excel = pd.ExcelFile(caminho)
    aba, observacao_aba = escolher_aba(arquivo_excel.sheet_names)
    linha_cabecalho, observacao_cabecalho = detectar_linha_cabecalho(caminho, aba)

    dados = pd.read_excel(caminho, sheet_name=aba, header=linha_cabecalho, dtype=str)
    metadados = {
        "aba_utilizada": aba,
        "linha_cabecalho": str(linha_cabecalho + 1),
        "quantidade_linhas": str(len(dados)),
        "quantidade_colunas": str(len(dados.columns)),
        "observacao": f"{observacao_aba} {observacao_cabecalho}",
    }
    return dados, metadados


def gerar_relatorio_ingestao(
    caminho_inventario: Path = INVENTARIO_ARQUIVOS,
    caminho_saida: Path = RELATORIO_INGESTAO,
) -> list[dict[str, str]]:
    registros = carregar_inventario(caminho_inventario)
    relatorio = []

    for registro in registros:
        if registro.get("status") != "ok":
            continue

        caminho_arquivo = caminho_absoluto(registro["caminho"])
        linha_relatorio = {
            "ano": registro.get("ano", ""),
            "nome_arquivo": registro.get("nome_arquivo", ""),
            "caminho": registro.get("caminho", ""),
            "aba_utilizada": "",
            "linha_cabecalho": "",
            "quantidade_linhas": "",
            "quantidade_colunas": "",
            "status": "ok",
            "observacao": "",
        }

        try:
            metadados = carregar_metadados_planilha(caminho_arquivo)
            linha_relatorio.update(metadados)
        except Exception as erro:
            linha_relatorio["status"] = "erro"
            linha_relatorio["observacao"] = f"Falha ao abrir arquivo: {erro}"

        relatorio.append(linha_relatorio)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    colunas = [
        "ano",
        "nome_arquivo",
        "caminho",
        "aba_utilizada",
        "linha_cabecalho",
        "quantidade_linhas",
        "quantidade_colunas",
        "status",
        "observacao",
    ]
    with caminho_saida.open("w", newline="", encoding="utf-8") as arquivo_csv:
        escritor = csv.DictWriter(arquivo_csv, fieldnames=colunas)
        escritor.writeheader()
        escritor.writerows(relatorio)

    return relatorio
